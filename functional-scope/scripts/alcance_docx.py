"""Constructor de documentos de Alcance Funcional sobre la plantilla corporativa de Campuslands.

Parte de `recursos/plantilla/Plantilla Word - Campuslands.docx` (membrete como imagen de fondo en el
encabezado de sección, márgenes y pie con número de página) y le aplica:
  * estilos del documento (cuerpo justificado, jerarquía de títulos tipo APA, tablas APA, estados),
  * la portada estándar,
  * control documental y tabla de contenido.

Uso (desde `functional-scope/`):
    python scripts/alcance_docx.py plantilla     # plantillas/Alcance Funcional - Plantilla.docx
    python scripts/alcance_docx.py muestra       # plantillas/Alcance Funcional - Muestra de formato.docx

La tabla de contenido se genera como campo de Word: al abrir el .docx, clic derecho → Actualizar campo (o Ctrl+A, F9).
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

RAIZ = Path(__file__).resolve().parent.parent
PLANTILLA_BASE = RAIZ / "recursos" / "plantilla" / "Plantilla Word - Campuslands.docx"

# Paleta tomada del membrete de la plantilla (azul y verde de la imagen de fondo).
AZUL = "0D497B"
VERDE = "5DAC6D"
VERDE_OSCURO = "2F7D45"   # verde con contraste suficiente para texto pequeño sobre blanco
TINTA = "262626"
GRIS = "595959"
TINTE = "EEF3F8"         # fondo suave de tarjetas y encabezados de tabla
REGLA = "D5DEE8"         # líneas finas entre filas
AMBAR = "B45309"

ANCHO_TEXTO = 8838        # twips: 12240 - 2 x 1701 de margen (el de la plantilla)
FUENTE = "Calibri"

PENDIENTE_TXT = "[PENDIENTE DE VALIDACIÓN]"
_MARCAS = re.compile(r"\[PENDIENTE DE VALIDACIÓN\]|CONFIRMADO|PROPUESTO|PENDIENTE")
_ESTILO_MARCA = {
    "[PENDIENTE DE VALIDACIÓN]": "Pendiente",
    "CONFIRMADO": "Estado Confirmado",
    "PROPUESTO": "Estado Propuesto",
    "PENDIENTE": "Estado Pendiente",
}

# ---------------------------------------------------------------- utilidades XML

_PPR_DESPUES_DE_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
    "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
    "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)
_RPR_DESPUES_DE_COLOR = (
    "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr",
    "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
    "w:specVanish", "w:oMath",
)
_RPR_DESPUES_DE_SHD = (
    "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
    "w:specVanish", "w:oMath",
)
_TCPR_ORDEN = ["w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd",
               "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"]
_TBLPR_ORDEN = ["w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
                "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd",
                "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"]


def _el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k.replace("_", ":", 1)), str(v))
    return e


def _insertar_ordenado(padre, hijo, orden: list[str]) -> None:
    """Inserta `hijo` en `padre` respetando el orden del esquema (Word es estricto con tcPr/tblPr)."""
    for previo in padre.findall(hijo.tag):
        padre.remove(previo)
    idx = orden.index(next(t for t in orden if qn(t) == hijo.tag))
    posteriores = {qn(t) for t in orden[idx + 1:]}
    for i, existente in enumerate(list(padre)):
        if existente.tag in posteriores:
            padre.insert(i, hijo)
            return
    padre.append(hijo)


def _borde_parrafo(estilo_o_par, lado: str, color: str, sz: int, espacio: int = 4) -> None:
    ppr = (estilo_o_par.element.get_or_add_pPr() if hasattr(estilo_o_par, "element")
           else estilo_o_par._p.get_or_add_pPr())
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.insert_element_before(pbdr, *_PPR_DESPUES_DE_PBDR)
    pbdr.append(_el(f"w:{lado}", w_val="single", w_sz=sz, w_space=espacio, w_color=color))


def _espaciado_letras(estilo, valor_twips: int) -> None:
    rpr = estilo.element.get_or_add_rPr()
    rpr.insert_element_before(_el("w:spacing", w_val=valor_twips), *_RPR_DESPUES_DE_COLOR)


def _sombrear_estilo_car(estilo, relleno: str) -> None:
    rpr = estilo.element.get_or_add_rPr()
    rpr.insert_element_before(_el("w:shd", w_val="clear", w_color="auto", w_fill=relleno),
                              *_RPR_DESPUES_DE_SHD)


def _color(rgb: str) -> RGBColor:
    return RGBColor.from_string(rgb)


def _campo(par, instruccion: str, resultado: str, estilo_car: str | None = None) -> None:
    """Campo complejo de Word (PAGE, SEQ, TOC…) con su resultado en caché."""
    def run():
        r = par.add_run()
        if estilo_car:
            r.style = estilo_car
        return r
    run()._r.append(_el("w:fldChar", w_fldCharType="begin"))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instruccion} "
    run()._r.append(it)
    run()._r.append(_el("w:fldChar", w_fldCharType="separate"))
    r = run()
    r.text = resultado
    run()._r.append(_el("w:fldChar", w_fldCharType="end"))


# ---------------------------------------------------------------- estilos

def _fuente(estilo, tam: float, color: str = TINTA, negrita=False, cursiva=False, mayus=False) -> None:
    f = estilo.font
    f.name = FUENTE
    f.size = Pt(tam)
    f.bold = negrita
    f.italic = cursiva
    f.all_caps = mayus or None
    f.color.rgb = _color(color)
    rpr = estilo.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    for a in ("w:eastAsia", "w:cs"):
        rfonts.set(qn(a), FUENTE)


def _parrafo(estilo, *, alin=WD_ALIGN_PARAGRAPH.LEFT, antes=0, despues=0, interlineado=1.15,
             mantener_sig=False, mantener_junto=False, sangria_izq=None, sangria_fr=None) -> None:
    pf = estilo.paragraph_format
    pf.alignment = alin
    pf.space_before = Pt(antes)
    pf.space_after = Pt(despues)
    pf.line_spacing = interlineado
    pf.keep_with_next = mantener_sig
    pf.keep_together = mantener_junto
    pf.widow_control = True
    if sangria_izq is not None:
        pf.left_indent = Twips(sangria_izq)
    if sangria_fr is not None:
        pf.first_line_indent = Twips(sangria_fr)


def _nuevo(doc, nombre: str, tipo=WD_STYLE_TYPE.PARAGRAPH, base: str | None = "Normal"):
    estilos = doc.styles
    if nombre in [s.name for s in estilos]:
        return estilos[nombre]
    e = estilos.add_style(nombre, tipo)
    if base and tipo == WD_STYLE_TYPE.PARAGRAPH:
        e.base_style = estilos[base]
    e.quick_style = True
    return e


def aplicar_estilos(doc) -> None:
    S = doc.styles

    # Cuerpo: Calibri 11, justificado, interlineado 1,15 y 6 pt posteriores (APA adaptado a documento comercial).
    n = S["Normal"]
    _fuente(n, 11, TINTA)
    _parrafo(n, alin=WD_ALIGN_PARAGRAPH.JUSTIFY, despues=6, interlineado=1.15)

    # Jerarquía de títulos (APA 7: negrita → negrita cursiva), alineados a la izquierda y nunca justificados.
    h1 = S["Heading 1"]
    _fuente(h1, 16, AZUL, negrita=True)
    _parrafo(h1, antes=20, despues=8, interlineado=1.05, mantener_sig=True, mantener_junto=True)
    _borde_parrafo(h1, "bottom", VERDE, 8, 3)
    h2 = S["Heading 2"]
    _fuente(h2, 13, AZUL, negrita=True)
    _parrafo(h2, antes=14, despues=5, interlineado=1.05, mantener_sig=True, mantener_junto=True)
    h3 = S["Heading 3"]
    _fuente(h3, 11.5, TINTA, negrita=True, cursiva=True)
    _parrafo(h3, antes=10, despues=4, interlineado=1.05, mantener_sig=True, mantener_junto=True)
    h4 = S["Heading 4"]
    _fuente(h4, 11, GRIS, negrita=True)
    _parrafo(h4, antes=8, despues=3, interlineado=1.05, mantener_sig=True, mantener_junto=True)

    # Listas con viñeta (numId 2 de la plantilla: «•» discreto).
    v = _nuevo(doc, "Viñeta")
    _fuente(v, 11, TINTA)
    _parrafo(v, alin=WD_ALIGN_PARAGRAPH.JUSTIFY, despues=3, interlineado=1.15)
    ppr = v.element.get_or_add_pPr()
    num = OxmlElement("w:numPr")
    num.append(_el("w:ilvl", w_val=0))
    num.append(_el("w:numId", w_val=2))
    ppr.insert_element_before(num, "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:spacing",
                              "w:ind", "w:jc", "w:outlineLvl")

    # Portada.
    t = _nuevo(doc, "Portada Tipo")
    _fuente(t, 11, VERDE_OSCURO, negrita=True, mayus=True)
    _espaciado_letras(t, 50)
    _parrafo(t, alin=WD_ALIGN_PARAGRAPH.CENTER, antes=70, despues=14)
    t = _nuevo(doc, "Portada Titulo")
    _fuente(t, 27, AZUL, negrita=True)
    _parrafo(t, alin=WD_ALIGN_PARAGRAPH.CENTER, despues=0, interlineado=1.08, sangria_izq=200)
    t.paragraph_format.right_indent = Twips(200)
    t = _nuevo(doc, "Portada Regla")
    _fuente(t, 2, VERDE)
    _parrafo(t, alin=WD_ALIGN_PARAGRAPH.CENTER, antes=20, despues=22, interlineado=1.0)
    t.paragraph_format.left_indent = Twips(3500)
    t.paragraph_format.right_indent = Twips(3500)
    _borde_parrafo(t, "bottom", VERDE, 18, 1)
    t = _nuevo(doc, "Portada Cliente")
    _fuente(t, 16, TINTA, negrita=True)
    _parrafo(t, alin=WD_ALIGN_PARAGRAPH.CENTER, despues=3, interlineado=1.05)
    t = _nuevo(doc, "Portada Area")
    _fuente(t, 12.5, GRIS)
    _parrafo(t, alin=WD_ALIGN_PARAGRAPH.CENTER, despues=0, interlineado=1.05)
    t = _nuevo(doc, "Portada Etiqueta")
    _fuente(t, 8.5, VERDE_OSCURO, negrita=True, mayus=True)
    _espaciado_letras(t, 30)
    _parrafo(t, alin=WD_ALIGN_PARAGRAPH.CENTER, antes=0, despues=2, interlineado=1.0, mantener_sig=True)
    t = _nuevo(doc, "Portada Valor")
    _fuente(t, 11, TINTA, negrita=True)
    _parrafo(t, alin=WD_ALIGN_PARAGRAPH.CENTER, despues=0, interlineado=1.05)
    t = _nuevo(doc, "Portada Nota")
    _fuente(t, 9.5, GRIS, cursiva=True)
    _parrafo(t, alin=WD_ALIGN_PARAGRAPH.CENTER, antes=0, despues=0, interlineado=1.1)
    t = _nuevo(doc, "Salto")
    _fuente(t, 1, TINTA)
    _parrafo(t, antes=0, despues=0, interlineado=1.0)
    t.paragraph_format.line_spacing = Pt(1)

    # Título sin numeración (no entra en la tabla de contenido).
    t = _nuevo(doc, "Título Sin Numerar")
    _fuente(t, 16, AZUL, negrita=True)
    _parrafo(t, antes=0, despues=10, interlineado=1.05, mantener_sig=True)
    _borde_parrafo(t, "bottom", VERDE, 8, 3)

    # Tabla de contenido.
    for i, (tam, neg, ant) in enumerate([(11, True, 8), (11, False, 2), (10.5, False, 0)], start=1):
        t = _nuevo(doc, f"toc {i}")
        _fuente(t, tam, TINTA if i > 1 else AZUL, negrita=neg)
        _parrafo(t, antes=ant, despues=2, interlineado=1.1, sangria_izq=(i - 1) * 300)
        t.paragraph_format.tab_stops.add_tab_stop(Twips(ANCHO_TEXTO), WD_TAB_ALIGNMENT.RIGHT,
                                                  WD_TAB_LEADER.DOTS)

    # Tablas (APA): número en negrita, título en cursiva sobre la tabla, nota debajo.
    t = _nuevo(doc, "Tabla Número")
    _fuente(t, 10.5, TINTA, negrita=True)
    _parrafo(t, antes=12, despues=0, interlineado=1.0, mantener_sig=True)
    t = _nuevo(doc, "Tabla Título")
    _fuente(t, 10.5, TINTA, cursiva=True)
    _parrafo(t, antes=0, despues=5, interlineado=1.0, mantener_sig=True)
    t = _nuevo(doc, "Tabla Texto")
    _fuente(t, 9.5, TINTA)
    _parrafo(t, antes=0, despues=0, interlineado=1.08)
    t = _nuevo(doc, "Tabla Encabezado")
    _fuente(t, 9.5, AZUL, negrita=True)
    _parrafo(t, antes=0, despues=0, interlineado=1.08, mantener_sig=True)
    t = _nuevo(doc, "Tabla Nota")
    _fuente(t, 9, GRIS)
    _parrafo(t, alin=WD_ALIGN_PARAGRAPH.JUSTIFY, antes=4, despues=12, interlineado=1.08)

    # Referencias (APA: sangría francesa de 1,27 cm, alineadas a la izquierda).
    t = _nuevo(doc, "Referencia")
    _fuente(t, 11, TINTA)
    _parrafo(t, despues=6, interlineado=1.15, sangria_izq=720, sangria_fr=-720)

    # Marcas de estado dentro del texto.
    c = _nuevo(doc, "Pendiente", WD_STYLE_TYPE.CHARACTER)
    c.font.bold = True
    c.font.color.rgb = _color(AMBAR)
    _sombrear_estilo_car(c, "FFF1CC")
    for nombre, color in [("Estado Confirmado", VERDE_OSCURO), ("Estado Propuesto", AZUL),
                          ("Estado Pendiente", AMBAR)]:
        c = _nuevo(doc, nombre, WD_STYLE_TYPE.CHARACTER)
        c.font.bold = True
        c.font.color.rgb = _color(color)


# ---------------------------------------------------------------- constructor

class Alcance:
    """Documento de alcance funcional sobre la plantilla corporativa."""

    def __init__(self, base: Path = PLANTILLA_BASE):
        self.doc = Document(str(base))
        cuerpo = self.doc.element.body
        for hijo in list(cuerpo):
            if hijo.tag != qn("w:sectPr"):
                cuerpo.remove(hijo)
        aplicar_estilos(self.doc)
        self._n_tabla = 0

    # -- texto
    def _texto(self, par, texto: str) -> None:
        """Agrega texto resaltando [PENDIENTE DE VALIDACIÓN] y los estados CONFIRMADO/PROPUESTO/PENDIENTE."""
        pos = 0
        for m in _MARCAS.finditer(texto):
            if m.start() > pos:
                par.add_run(texto[pos:m.start()])
            par.add_run(m.group(0)).style = _ESTILO_MARCA[m.group(0)]
            pos = m.end()
        if pos < len(texto):
            par.add_run(texto[pos:])

    def parrafo(self, texto: str, estilo: str | None = None):
        p = self.doc.add_paragraph(style=estilo)
        self._texto(p, texto)
        return p

    def h1(self, t, nueva_pagina: bool = False):
        p = self.doc.add_paragraph(t, style="Heading 1")
        p.paragraph_format.page_break_before = nueva_pagina or None
        return p
    def h2(self, t): return self.doc.add_paragraph(t, style="Heading 2")
    def h3(self, t): return self.doc.add_paragraph(t, style="Heading 3")
    def h4(self, t): return self.doc.add_paragraph(t, style="Heading 4")

    def vinetas(self, items: list[str]) -> None:
        for it in items:
            self.parrafo(it, "Viñeta")

    def referencias(self, items: list[str]) -> None:
        for it in sorted(items, key=str.lower):
            self.parrafo(it, "Referencia")

    def salto_pagina(self) -> None:
        p = self.doc.add_paragraph(style="Salto")
        p.add_run().add_break(WD_BREAK.PAGE)

    # -- tablas
    def _tabla(self, filas: int, anchos: list[int], ancho_total: int | None = None, centrada=False):
        cols = len(anchos)
        t = self.doc.add_table(rows=filas, cols=cols)
        t.autofit = False
        tblpr = t._tbl.tblPr
        total = ancho_total or sum(anchos)
        _insertar_ordenado(tblpr, _el("w:tblW", w_w=total, w_type="dxa"), _TBLPR_ORDEN)
        _insertar_ordenado(tblpr, _el("w:tblLayout", w_type="fixed"), _TBLPR_ORDEN)
        if centrada:
            _insertar_ordenado(tblpr, _el("w:jc", w_val="center"), _TBLPR_ORDEN)
        mar = _el("w:tblCellMar")
        for lado, v in (("top", 70), ("left", 110), ("bottom", 70), ("right", 110)):
            mar.append(_el(f"w:{lado}", w_w=v, w_type="dxa"))
        _insertar_ordenado(tblpr, mar, _TBLPR_ORDEN)
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), anchos):
            gc.set(qn("w:w"), str(w))
        for fila in t.rows:
            for c, w in zip(fila.cells, anchos):
                _insertar_ordenado(c._tc.get_or_add_tcPr(), _el("w:tcW", w_w=w, w_type="dxa"), _TCPR_ORDEN)
        return t

    @staticmethod
    def _celda(c, relleno: str | None = None, v_centro=False, bordes: dict | None = None) -> None:
        tcpr = c._tc.get_or_add_tcPr()
        if bordes:
            b = OxmlElement("w:tcBorders")
            for lado in ("top", "left", "bottom", "right"):
                if lado in bordes:
                    color, sz = bordes[lado]
                    b.append(_el(f"w:{lado}", w_val="single" if color else "nil", w_sz=sz,
                                 w_space=0, w_color=color or "auto"))
            _insertar_ordenado(tcpr, b, _TCPR_ORDEN)
        if relleno:
            _insertar_ordenado(tcpr, _el("w:shd", w_val="clear", w_color="auto", w_fill=relleno), _TCPR_ORDEN)
        if v_centro:
            _insertar_ordenado(tcpr, _el("w:vAlign", w_val="center"), _TCPR_ORDEN)

    @staticmethod
    def _fila(fila, encabezado=False, no_partir=True) -> None:
        trpr = fila._tr.get_or_add_trPr()
        if no_partir:
            trpr.append(_el("w:cantSplit"))
        if encabezado:
            trpr.append(_el("w:tblHeader"))

    def _rellenar(self, celda, texto: str, estilo: str) -> None:
        p = celda.paragraphs[0]
        p.style = estilo
        self._texto(p, texto)

    def tabla(self, titulo: str, encabezados: list[str], filas: list[list[str]], pesos: list[float],
              nota: str | None = None) -> None:
        """Tabla en formato APA: «Tabla N» en negrita, título en cursiva, reglas horizontales, nota debajo."""
        self._n_tabla += 1
        p = self.doc.add_paragraph(style="Tabla Número")
        p.add_run("Tabla ")
        _campo(p, "SEQ Tabla \\* ARABIC", str(self._n_tabla))
        self.doc.add_paragraph(titulo, style="Tabla Título")

        anchos = [round(ANCHO_TEXTO * w / sum(pesos)) for w in pesos]
        anchos[-1] += ANCHO_TEXTO - sum(anchos)
        t = self._tabla(1 + len(filas), anchos)
        for i, h in enumerate(encabezados):
            c = t.rows[0].cells[i]
            self._rellenar(c, h, "Tabla Encabezado")
            self._celda(c, TINTE, v_centro=True,
                        bordes={"top": (AZUL, 10), "bottom": (AZUL, 6)})
        self._fila(t.rows[0], encabezado=True)
        for r, datos in enumerate(filas, start=1):
            ultima = r == len(filas)
            for i, valor in enumerate(datos):
                c = t.rows[r].cells[i]
                self._rellenar(c, valor, "Tabla Texto")
                self._celda(c, bordes={"bottom": (AZUL, 10) if ultima else (REGLA, 4)})
            self._fila(t.rows[r])
        if nota:
            p = self.doc.add_paragraph(style="Tabla Nota")
            r = p.add_run("Nota. ")
            r.italic = True
            self._texto(p, nota)
        else:
            self.doc.add_paragraph(style="Tabla Nota")

    def ficha_funcional(self, rf_id: str, nombre: str, campos: list[tuple[str, str]]) -> None:
        """Ficha uniforme por funcionalidad (manual §5.1): encabezado «RF-001 · Nombre» y filas campo/contenido."""
        anchos = [2150, ANCHO_TEXTO - 2150]
        t = self._tabla(1 + len(campos), anchos)
        cab = t.rows[0].cells[0].merge(t.rows[0].cells[1])
        p = cab.paragraphs[0]
        p.style = "Tabla Encabezado"
        r = p.add_run(f"{rf_id}  ·  {nombre}")
        r.font.color.rgb = _color("FFFFFF")
        r.font.size = Pt(10.5)
        self._celda(cab, AZUL, v_centro=True)
        self._fila(t.rows[0])
        t.rows[0].cells[0].paragraphs[0].paragraph_format.keep_with_next = True
        for i, (etq, val) in enumerate(campos, start=1):
            ultima = i == len(campos)
            c0, c1 = t.rows[i].cells
            self._rellenar(c0, etq, "Tabla Encabezado")
            self._celda(c0, TINTE, bordes={"bottom": (REGLA if not ultima else AZUL, 4 if not ultima else 10)})
            self._rellenar(c1, val, "Tabla Texto")
            self._celda(c1, bordes={"bottom": (REGLA if not ultima else AZUL, 4 if not ultima else 10)})
            self._fila(t.rows[i])
            if not ultima:
                for c in (c0, c1):
                    c.paragraphs[0].paragraph_format.keep_with_next = True
        self.doc.add_paragraph(style="Tabla Nota")

    # -- portada (diseño definitivo)
    def portada(self, d: dict) -> None:
        """Portada estándar. Campos: titulo, cliente (obligatorios); tipo, area, proveedor, version, fecha,
        clasificacion, estado, nota (con valores por defecto)."""
        tipo = d.get("tipo", "Documento de alcance funcional")
        self.doc.add_paragraph(tipo, style="Portada Tipo")
        self.doc.add_paragraph(d["titulo"], style="Portada Titulo")
        self.doc.add_paragraph(style="Portada Regla")
        self.doc.add_paragraph(d["cliente"], style="Portada Cliente")
        if d.get("area"):
            self.doc.add_paragraph(d["area"], style="Portada Area")

        celdas = [
            ("Preparado para", d["cliente"]),
            ("Preparado por", d.get("proveedor", "Campuslands S.A.S. BIC · Full Service")),
            ("Versión", d.get("version", PENDIENTE_TXT)),
            ("Fecha", d.get("fecha", PENDIENTE_TXT)),
            ("Clasificación", d.get("clasificacion", "Confidencial")),
            ("Estado del documento", d.get("estado", "Borrador")),
        ]
        sep = self.doc.add_paragraph(style="Salto")
        sep.paragraph_format.space_before = Pt(0)
        sep.paragraph_format.space_after = Pt(0)
        sep.paragraph_format.line_spacing = Pt(1)
        sep.paragraph_format.space_before = Pt(50)

        ancho = 7800
        t = self._tabla(3, [ancho // 2, ancho // 2], centrada=True)
        blanco = ("FFFFFF", 24)
        for i, (etq, val) in enumerate(celdas):
            c = t.rows[i // 2].cells[i % 2]
            self._celda(c, TINTE, v_centro=True,
                        bordes={"top": blanco, "left": blanco, "bottom": blanco, "right": blanco})
            c.paragraphs[0].style = "Portada Etiqueta"
            c.paragraphs[0].add_run(etq)
            v = c.add_paragraph(style="Portada Valor")
            self._texto(v, val)
            # Aire interno de la tarjeta.
            _insertar_ordenado(c._tc.get_or_add_tcPr(), self._mar(130, 130), _TCPR_ORDEN)
        for fila in t.rows:
            self._fila(fila)

        nota = d.get("nota", "Documento confidencial — Para revisión técnica, comercial y jurídica")
        p = self.doc.add_paragraph(nota, style="Portada Nota")
        p.paragraph_format.space_before = Pt(56)

    @staticmethod
    def _mar(arriba: int, abajo: int):
        m = _el("w:tcMar")
        m.append(_el("w:top", w_w=arriba, w_type="dxa"))
        m.append(_el("w:bottom", w_w=abajo, w_type="dxa"))
        return m

    # -- control documental y contenido
    def control_documental(self, d: dict) -> None:
        self.h1("Control documental", nueva_pagina=True)
        self.h2("Responsables")
        pend = PENDIENTE_TXT
        self.tabla(
            "Responsables y revisiones del documento",
            ["Rol", "Nombre", "Cargo o área", "Revisión"],
            [
                ["Autor del alcance", d.get("autor", pend), "Campuslands Full Service", "Redacción"],
                ["Líder técnico", pend, pend, "Técnica"],
                ["Responsable comercial", pend, pend, "Comercial"],
                ["Asesoría jurídica", pend, pend, "Jurídica"],
                ["Patrocinador (cliente)", pend, pend, "Validación del alcance"],
            ],
            [2.0, 2.6, 2.6, 1.6],
            nota="Una versión solo se declara apta para revisión final cuando no quedan pendientes críticos.",
        )
        self.h2("Historial de cambios")
        self.tabla(
            "Versiones del documento",
            ["Versión", "Fecha", "Autor", "Descripción del cambio", "Estado"],
            [[d.get("version", pend), d.get("fecha", pend), d.get("autor", pend), pend, d.get("estado", "Borrador")]],
            [0.9, 2.1, 2.1, 2.7, 1.2],
        )
        self.h2("Estados de la información")
        self.tabla(
            "Estados usados en este documento",
            ["Estado", "Significado"],
            [
                ["CONFIRMADO", "Respaldado por una fuente citada o por una aprobación expresa."],
                ["PROPUESTO", "Recomendación del equipo, aún sin aceptación del cliente."],
                ["PENDIENTE", "Dato o decisión por definir; se marca como " + PENDIENTE_TXT + "."],
            ],
            [1.6, 6.4],
        )

    def tabla_de_contenido(self) -> None:
        p = self.doc.add_paragraph("Tabla de contenido", style="Título Sin Numerar")
        p.paragraph_format.page_break_before = True
        p = self.doc.add_paragraph(style="Normal")
        _campo(p, 'TOC \\o "1-3" \\h \\z \\u',
               "Clic derecho sobre este texto → Actualizar campo, para generar la tabla de contenido.")

    def guardar(self, ruta: Path, titulo: str, tema: str = "Alcance funcional") -> Path:
        cp = self.doc.core_properties
        cp.title = titulo
        cp.subject = tema
        cp.author = "Campuslands Full Service"
        cp.last_modified_by = "Campuslands Full Service"
        cp.language = "es-CO"
        ruta.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(ruta))
        return ruta


# ---------------------------------------------------------------- documentos de ejemplo

def construir_plantilla(salida: Path) -> Path:
    a = Alcance()
    a.portada({
        "titulo": "[Nombre del proyecto]",
        "cliente": "[Razón social del cliente]",
        "area": "[Dependencia o unidad del cliente]",
        "version": "[0.1]",
        "fecha": "[Día de mes de año]",
    })
    a.control_documental({"version": "[0.1]", "fecha": "[Día de mes de año]", "autor": "[Nombre del autor]"})
    a.tabla_de_contenido()
    return a.guardar(salida, "Alcance Funcional — Plantilla")


def construir_muestra(salida: Path) -> Path:
    """Portada del proyecto de referencia (contenido de `Portada.docx`) + páginas que muestran cada estilo."""
    a = Alcance()
    datos = {
        "titulo": "Servicio inteligente de cobro persuasivo y fiscalización tributaria con IA",
        "cliente": "Alcaldía Municipal de Girón",
        "area": "Industria y Comercio (ICA)",
        "version": "2.0",
        "fecha": "28 de octubre de 2026",
        "autor": PENDIENTE_TXT,
        "nota": "Documento confidencial — Para revisión técnica, jurídica, financiera y contractual",
    }
    a.portada(datos)
    a.control_documental(datos)
    a.tabla_de_contenido()

    a.h1("1. Resumen ejecutivo", nueva_pagina=True)
    a.parrafo("Este párrafo de muestra ilustra el formato del cuerpo del texto: Calibri 11 pt, interlineado de "
              "1,15, seis puntos de espacio posterior y alineación justificada, de modo que el bloque de texto "
              "conserve márgenes limpios a izquierda y derecha. El contenido de esta página es solo de ejemplo "
              "para evaluar el formato y no describe compromisos del proyecto.")
    a.parrafo("Los datos aún no confirmados se marcan así: " + PENDIENTE_TXT + ". Los estados de una "
              "funcionalidad se resaltan como CONFIRMADO, PROPUESTO o PENDIENTE.")
    a.h2("1.1 Subtítulo de segundo nivel")
    a.parrafo("Los títulos siguen la jerarquía APA: primer nivel en negrita con línea verde, segundo nivel en "
              "negrita, tercer nivel en negrita cursiva. Todos van alineados a la izquierda y se mantienen "
              "unidos al párrafo que sigue.")
    a.h3("1.1.1 Subtítulo de tercer nivel")
    a.vinetas([
        "Las viñetas conservan la justificación y una separación de tres puntos entre elementos.",
        "Cada elemento describe un resultado verificable, no un verbo amplio como gestionar o analizar.",
        "Las funcionalidades de fases futuras nunca se listan como entregables actuales.",
    ])
    a.tabla("Ejemplo de tabla en formato APA",
            ["Módulo", "Descripción", "Prioridad", "Estado"],
            [["MOD-01", "Descripción breve y verificable del módulo de ejemplo.", "Alta", "PROPUESTO"],
             ["MOD-02", "Segundo módulo de ejemplo con un texto algo más largo para comprobar el ajuste "
                        "de línea dentro de la celda.", "Media", "PENDIENTE"],
             ["MOD-03", "Tercer módulo de ejemplo.", "Baja", "CONFIRMADO"]],
            [1.0, 4.6, 1.2, 1.6],
            nota="Las tablas llevan número y título sobre ellas, reglas horizontales y la fuente o aclaración "
                 "debajo, como indica APA.")
    a.h2("1.2 Ficha funcional")
    a.ficha_funcional("RF-001", "Nombre claro de la funcionalidad", [
        ("Objetivo y usuario", "Necesidad, rol y valor. " + PENDIENTE_TXT),
        ("Precondiciones y entradas", "Datos, accesos y evento inicial."),
        ("Flujo principal", "1. Paso observable del usuario. 2. Respuesta del sistema."),
        ("Reglas y excepciones", "Cálculos, validaciones, errores y casos límite."),
        ("Salidas", "Pantallas, reportes, notificaciones o acciones."),
        ("Dependencias", "Sistemas, información, decisiones y terceros."),
        ("Criterios de aceptación", "Pruebas observables y medibles."),
        ("Estado", "PROPUESTO — fuente: [documento, sección, fecha]"),
    ])
    a.h1("Referencias")
    a.referencias([
        "Apellido, N. (2026). Título de la fuente en cursiva. Editorial. https://ejemplo.org",
        "Entidad. (2026, 28 de octubre). Título del documento [Acta de reunión]. Campuslands.",
    ])
    return a.guardar(salida, "Alcance Funcional — Muestra de formato")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    cmd = sys.argv[1] if len(sys.argv) > 1 else "plantilla"
    destino = RAIZ / "plantillas"
    if cmd == "plantilla":
        print(construir_plantilla(destino / "Alcance Funcional - Plantilla.docx"))
    elif cmd == "muestra":
        print(construir_muestra(destino / "Alcance Funcional - Muestra de formato.docx"))
    else:
        sys.exit("Uso: alcance_docx.py [plantilla|muestra]")
